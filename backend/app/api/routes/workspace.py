import copy
import hashlib
import io
import json
import math
import tempfile
import uuid
from dataclasses import asdict
from datetime import date
from typing import Annotated
from pathlib import Path
from urllib.parse import quote
from fastapi import APIRouter, HTTPException, UploadFile, File, Form, Request, Depends
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from app.api.deps import CurrentUser, SessionDep
from app.api.routes.projects import ProjectDep, EditProject, require_features, save_project, public
from app.project_models import (
    ProjectRecord,
    DataWrite,
    DatabaseWrite,
    DatabaseDate,
    Scope,
    Ask,
    ChartPreference,
    ChartSettings,
    LevelAdjustmentApplyRequest,
    TwinSceneWrite,
)
from app.database_view import replace_date_rows
from app.charts import chart_data
from app.chart_links import make_chart_link
from app import monitoring as m
from app.level_adjustment import (
    LevelAdjustmentError,
    cross_check_level_outputs,
    parse_level_adjustment_files,
    solve_level_network,
)
from app.core.config import settings
from app.level_imports import prepare_level_import, generate_level_reports
from app.twin_dwg import TwinDwgError, scene_from_dwg

router=APIRouter(prefix="/projects/{project_id}",tags=["workspace"])
MAX_UPLOAD=20*1024*1024
TWIN_POINT_MIN = -2.0
TWIN_POINT_MAX = 2.0
DatabaseProject = Annotated[ProjectRecord, Depends(require_features('database_data'))]
ChartProject = Annotated[ProjectRecord, Depends(require_features('database_charts'))]
EntryProject = Annotated[ProjectRecord, Depends(require_features('data_entry'))]
AnalysisProject = Annotated[ProjectRecord, Depends(require_features('analysis'))]
ReportProject = Annotated[ProjectRecord, Depends(require_features('reports'))]
CatalogProject = Annotated[ProjectRecord, Depends(require_features('data_entry', 'database_data', 'database_charts', 'analysis', 'reports'))]

@router.get("/chart-link")
def chart_link(measurement_key: str, request: Request, project: ChartProject):
    reader = m.MonitoringReader(m.service_for(project))
    if measurement_key not in {item["key"] for item in reader.list_measurements()}:
        raise HTTPException(422, "当前工程中不存在此测项")
    return make_chart_link(request, project.id, measurement_key)

@router.get("/charts")
def charts(measurement_key: str, project: ChartProject, current_user: CurrentUser, session: SessionDep):
    result = measurement_view(measurement_key, project, current_user, session)
    # Chart-only accounts must not receive raw readings or report table fields.
    return {key: value for key, value in result.items()
            if key not in {'rows', 'columns', 'raw_rows', 'raw_columns'}}

@router.get('/database')
def database(measurement_key: str, project: DatabaseProject, current_user: CurrentUser, session: SessionDep):
    result = measurement_view(measurement_key, project, current_user, session)
    fields = {'measurement_key', 'label', 'monitor', 'unit', 'dates', 'effective_date',
              'columns', 'rows', 'raw_rows', 'raw_columns', 'warnings'}
    return {key: value for key, value in result.items() if key in fields} | {
        'end_date': result['settings']['end_date'],
    }

def measurement_view(measurement_key, project, current_user, session):
    saved = session.get(ChartPreference, (project.id, current_user.id, measurement_key))
    preference = ChartSettings(measurement_key=measurement_key, **(saved.settings if saved else {}))
    try:
        return chart_data(project, preference)
    except (ValueError, KeyError, TypeError) as exc:
        data_error(exc)

@router.put('/database-date')
def database_date(body: DatabaseDate, project: DatabaseProject, current_user: CurrentUser, session: SessionDep):
    reader = m.MonitoringReader(m.service_for(project))
    try:
        dates = reader.list_dates(body.measurement_key)
        if body.end_date and body.end_date not in dates:
            raise ValueError("请选择当前测项中存在的监测日期")
    except (ValueError, KeyError, TypeError) as exc:
        data_error(exc)
    preference = session.get(ChartPreference, (project.id, current_user.id, body.measurement_key))
    if preference is None:
        preference = ChartPreference(project_id=project.id, user_id=current_user.id, measurement_key=body.measurement_key)
    preference.settings = {**preference.settings, 'end_date': body.end_date}
    session.add(preference)
    session.commit()
    return {'message': '数据库日期已保存'}

@router.put("/chart-settings")
def chart_settings(body: ChartSettings, project: ChartProject, current_user: CurrentUser, session: SessionDep):
    reader = m.MonitoringReader(m.service_for(project))
    try:
        available = set(reader.list_points(body.measurement_key))
        if any(point not in available for point in body.points):
            raise ValueError("所选点位不属于当前测项，请刷新后重新选择")
        if body.end_date and body.end_date not in reader.list_dates(body.measurement_key):
            raise ValueError("请选择当前测项中存在的监测日期")
    except (ValueError, KeyError, TypeError) as exc:
        data_error(exc)
    preference = session.get(ChartPreference, (project.id, current_user.id, body.measurement_key))
    if preference is None:
        preference = ChartPreference(project_id=project.id, user_id=current_user.id, measurement_key=body.measurement_key)
    preference.settings = body.model_dump(exclude={"measurement_key"})
    session.add(preference)
    session.commit()
    return {"message": "图表设置已保存"}

def data_error(exc):
    raise HTTPException(422,str(exc)) from None

def document_public(document):
    return {k:v for k,v in document.items() if k not in {"source_path","file_hash"}}

def bounded_upload(file):
    content=file.file.read(MAX_UPLOAD+1)
    if len(content)>MAX_UPLOAD:
        raise HTTPException(413,"文件超过 20 MB，请拆分后导入")
    if not content:
        raise HTTPException(422,"文件为空")
    return content


def validate_twin_scene(scene):
    if not isinstance(scene, dict) or scene.get("schema_version") != 1:
        raise ValueError("数字孪生场景格式无效或版本不支持")
    # Sensor recognition belongs to the DWG import pipeline and is deliberately
    # discarded here. The web scene stores only structural geometry plus points
    # explicitly selected from the project's database catalog.
    scene = copy.deepcopy(scene)
    scene.pop("sensors", None)
    scene.pop("sensor_binding_rules", None)
    scene.pop("recognition", None)
    scene.setdefault("placed_points", [])
    layers = scene.get("layers", [])
    objects = scene.get("objects", [])
    placed_points = scene.get("placed_points", [])
    if not isinstance(layers, list) or not isinstance(objects, list) or not isinstance(placed_points, list):
        raise ValueError("数字孪生图层、对象和已选测点必须是数组")
    if len(layers) > 500 or len(objects) > 20000 or len(placed_points) > 20000:
        raise ValueError("数字孪生场景对象数量超过限制")
    layer_ids = set()
    for layer in layers:
        if not isinstance(layer, dict) or not isinstance(layer.get("id"), str) or not isinstance(layer.get("name"), str):
            raise ValueError("数字孪生图层格式无效")
        if layer["id"] in layer_ids:
            raise ValueError("数字孪生图层 ID 重复")
        layer_ids.add(layer["id"])
    for obj in objects:
        if not isinstance(obj, dict) or not isinstance(obj.get("id"), str) or not isinstance(obj.get("geometry"), dict):
            raise ValueError("数字孪生几何对象格式无效")
        if obj.get("layer_id") and obj["layer_id"] not in layer_ids:
            raise ValueError("数字孪生对象引用了不存在的图层")
    point_ids = set()
    for point in placed_points:
        if not isinstance(point, dict) or not isinstance(point.get("id"), str):
            raise ValueError("数字孪生已选测点格式无效")
        if not isinstance(point.get("measurement_key"), str) or not isinstance(point.get("point"), str):
            raise ValueError("数字孪生已选测点必须引用测项和点号")
        if point["id"] in point_ids:
            raise ValueError("数字孪生已选测点 ID 重复")
        point_ids.add(point["id"])
        position = point.get("position")
        if position is not None:
            if not isinstance(position, list) or len(position) != 2 or any(not isinstance(v, (int, float)) or not math.isfinite(v) or v < TWIN_POINT_MIN or v > TWIN_POINT_MAX for v in position):
                raise ValueError(f"测点位置必须在 {TWIN_POINT_MIN} 到 {TWIN_POINT_MAX} 的沙盘扩展坐标范围内")
    return scene


@router.get("/twin/scene")
def twin_scene(project: AnalysisProject):
    return {"revision": project.revision, "scene": project.payload.get("twin_scene")}


@router.post("/twin/import")
def import_twin_scene(project: EntryProject, session: SessionDep, file: UploadFile = File(...)):
    filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(filename).suffix.lower() != ".json":
        raise HTTPException(422, "数字孪生导入暂支持场景 JSON 文件")
    try:
        content = bounded_upload(file)
        scene = validate_twin_scene(json.loads(content.decode("utf-8-sig")))
        payload = copy.deepcopy(project.payload)
        payload["twin_scene"] = scene
        save_project(session, project, project.revision, payload=payload)
    except (UnicodeError, json.JSONDecodeError, ValueError, TypeError) as exc:
        data_error(exc)
    return {**public(project, session), "scene": scene, "message": "数字孪生场景已导入"}


@router.post("/twin/dwg-preview")
def preview_twin_dwg(project: EntryProject, file: UploadFile = File(...)):
    filename = Path((file.filename or "").replace("\\", "/")).name
    if Path(filename).suffix.lower() != ".dwg":
        raise HTTPException(422, "DWG 自动生成暂支持 .dwg 文件")
    try:
        content = bounded_upload(file)
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / filename
            path.write_bytes(content)
            scene = validate_twin_scene(scene_from_dwg(path, filename))
    except (TwinDwgError, UnicodeError, ValueError, TypeError, OSError) as exc:
        data_error(exc)
    scene["project"] = {"id": str(project.id), "name": project.name, "revision_at_export": project.revision}
    return {"revision": project.revision, "scene": scene, "message": "DWG 已解析为数字孪生场景，请核对后保存"}


def _save_twin_scene(body: TwinSceneWrite, project: EntryProject, session: SessionDep):
    try:
        scene = validate_twin_scene(body.scene)
        payload = copy.deepcopy(project.payload)
        payload["twin_scene"] = scene
        save_project(session, project, body.revision, payload=payload)
    except (ValueError, TypeError) as exc:
        data_error(exc)
    return {**public(project, session), "scene": scene, "message": "数字孪生场景已保存"}


@router.put("/twin/scene")
def save_twin_scene(body: TwinSceneWrite, project: EntryProject, session: SessionDep):
    return _save_twin_scene(body, project, session)


@router.post("/twin/scene")
def save_twin_scene_compat(body: TwinSceneWrite, project: EntryProject, session: SessionDep):
    # Compatibility for an already-open browser tab that still has the old
    # POST-based bundle. New clients use PUT, but both paths share validation
    # and optimistic revision checks.
    return _save_twin_scene(body, project, session)


LEVEL_ADJUSTMENT_EXTENSIONS = {".dat", ".bm1", ".hdf", ".in1", ".gco", ".ou1", ".our"}


def _level_adjustment_directory(project_id, import_id):
    return settings.STORAGE_DIR / str(project_id) / "level-adjustments" / import_id


def _level_adjustment_public(result, checks, parsed):
    return {
        "points": {point: asdict(value) for point, value in result.points.items()},
        "observations": [asdict(value) for value in result.observations],
        "corrections": result.corrections,
        "total_distance_km": result.total_distance_km,
        "closure_error_mm": result.closure_error_mm,
        "quality": asdict(result.quality),
        "cross_checks": [asdict(check) for check in checks],
        "known_points": parsed.known_points,
        "route_start": parsed.route_start,
        "route_end": parsed.route_end,
    }


@router.post("/level-adjustment/preview")
def preview_level_adjustment(
    project: EntryProject,
    files: list[UploadFile] = File(...),
    weight_mode: str = Form(...),
    observed_date: str = Form(...),
):
    if weight_mode not in {"distance", "station"}:
        raise HTTPException(422, "请选择按距离定权或按测站定权")
    if len(files) > 7:
        raise HTTPException(422, "一次最多导入 7 个原始文件")
    names: dict[str, bytes] = {}
    for upload in files:
        name = Path((upload.filename or "").replace("\\", "/")).name
        if Path(name).suffix.lower() not in LEVEL_ADJUSTMENT_EXTENSIONS:
            raise HTTPException(422, "仅支持 DAT、BM1、hdf、in1、gco、ou1、our 文件")
        content = bounded_upload(upload)
        names[name] = content
    import_id = uuid.uuid4().hex
    try:
        date.fromisoformat(observed_date)
        if len({Path(name).suffix.lower() for name in names}) != len(names):
            raise ValueError("同一后缀只能上传一个文件")
        mode, parsed, result, checks = prepare_level_import(names, weight_mode)
        target_points: list[str] = []
        try:
            target_points = m.MonitoringReader(m.service_for(project)).list_points("level:建筑物位移")
        except (ValueError, KeyError, TypeError):
            target_points = []
        mapping_candidates = {
            point: ([point] if not target_points or point in target_points else [candidate for candidate in target_points if point.casefold() in candidate.casefold()])
            for point in result.points
        }
        directory = _level_adjustment_directory(project.id, import_id)
        directory.mkdir(parents=True, exist_ok=False)
        for name, content in names.items():
            (directory / name).write_bytes(content)
        generated = {}
        if mode == "adjusted":
            dat_name = next((n for n in names if n.lower().endswith(".dat")), next(iter(names)))
            generated = generate_level_reports(dat_name, parsed, result, observed_date)
            output = directory / "generated"
            output.mkdir()
            for name, content in generated.items():
                (output / name).write_bytes(content)
        (directory / "manifest.json").write_text(
            json.dumps({"mode": mode, "weight_mode": weight_mode, "observed_date": observed_date, "files": list(names), "generated": list(generated)}, ensure_ascii=False),
            encoding="utf-8",
        )
    except (LevelAdjustmentError, ValueError, KeyError, TypeError, OSError) as exc:
        data_error(exc)
    return {
        "import_id": import_id,
        "mode": mode,
        "message": "已完成水准平差，并生成派生文件" if mode == "adjusted" else "未上传控制点文件，不做平差处理，直接使用观测高程进行点号映射",
        "generated_files": [{"name": name, "bytes": len(content), "url": f"/api/v1/projects/{project.id}/level-adjustment/{import_id}/files/{quote(name)}"} for name, content in generated.items()],
        "files": [{"name": name, "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()} for name, content in names.items()],
        "result": _level_adjustment_public(result, checks, parsed),
        "quality": asdict(result.quality),
        "cross_checks": [asdict(check) for check in checks],
        "mapping_candidates": mapping_candidates,
        "project_revision": project.revision,
    }


@router.get("/level-adjustment/{import_id}/files/{filename}")
def download_level_adjustment_file(import_id: str, filename: str, project: EntryProject):
    if len(import_id) != 32 or any(c not in "0123456789abcdef" for c in import_id):
        raise HTTPException(404, "导入预览不存在")
    directory = _level_adjustment_directory(project.id, import_id)
    manifest_path = directory / "manifest.json"
    if not manifest_path.is_file():
        raise HTTPException(404, "导入预览不存在")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if filename not in manifest.get("generated", []) or Path(filename).name != filename:
        raise HTTPException(404, "生成文件不存在")
    return Response((directory / "generated" / filename).read_bytes(), media_type="application/octet-stream",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


@router.put("/level-adjustment/apply")
def apply_level_adjustment(body: LevelAdjustmentApplyRequest, project: EntryProject, session: SessionDep):
    directory = _level_adjustment_directory(project.id, body.import_id)
    manifest_path = directory / "manifest.json"
    if not manifest_path.is_file():
        raise HTTPException(404, "导入预览不存在或已过期")
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("weight_mode") != body.weight_mode or manifest.get("observed_date") != body.observed_date:
            raise ValueError("应用参数与预览结果不一致，请重新预览")
        files = {name: (directory / name).read_bytes() for name in manifest.get("files", [])}
        mode, parsed, result, checks = prepare_level_import(files, body.weight_mode)
        if any(check.status == "mismatch" for check in checks):
            raise ValueError("原始派生文件与重新计算结果不一致，不能写入")
        if not result.quality.can_apply and not body.allow_out_of_tolerance:
            raise ValueError("闭合差或限差检查未通过；确认允许超限结果后才能写入")
        if not body.mappings:
            raise ValueError("至少选择一个点号映射")
        measurement_key = body.measurement_key
        if not measurement_key.startswith("level:"):
            raise ValueError("目标测项必须是水准测项")
        payload = copy.deepcopy(project.payload)
        measurements = payload.setdefault("measurements", {})
        rows = list(measurements.get(measurement_key, []))
        existing_points = {str(row.get("点号", "")).strip() for row in rows if row.get("点号")}
        enabled = [mapping for mapping in body.mappings if mapping.enabled]
        if not enabled:
            raise ValueError("至少选择一个点号映射")
        if len({x.database_point for x in enabled}) != len(enabled) or len({x.adjusted_point for x in enabled}) != len(enabled):
            raise ValueError("点号映射重复：每个数据库点号只能对应一个观测点")
        if any(mapping.adjusted_point not in result.points for mapping in enabled):
            raise ValueError("映射中包含不存在的平差点号")
        if existing_points and any(mapping.database_point not in existing_points for mapping in enabled):
            raise ValueError("映射中包含不属于目标水准测项的数据库点号")
        for mapping in enabled:
            value = result.points[mapping.adjusted_point].elevation
            matching = [row for row in rows if str(row.get("点号", "")).strip() == mapping.database_point and str(row.get("时间", "")) == body.observed_date]
            if matching:
                if len(matching) > 1:
                    raise ValueError(f"{mapping.database_point} 在 {body.observed_date} 存在重复记录")
                matching[0]["测量值"] = value
            else:
                rows.append({"时间": body.observed_date, "点号": mapping.database_point, "测量值": value, "备注": "原始水准平差导入" if mode == "adjusted" else "DAT 观测高程直接导入（未平差）"})
        if rows and any("序号" in row for row in rows):
            for sequence, row in enumerate(rows, start=1):
                row["序号"] = sequence
        payload["measurements"][measurement_key] = rows
        payload.setdefault("level_adjustment_imports", []).append({"import_id": body.import_id, "mode": mode, "observed_date": body.observed_date, "weight_mode": body.weight_mode if mode == "adjusted" else None, "mappings": [mapping.model_dump() for mapping in enabled]})
        save_project(session, project, body.revision, payload=payload)
    except HTTPException:
        raise
    except (LevelAdjustmentError, ValueError, KeyError, TypeError, OSError) as exc:
        data_error(exc)
    return {**public(project, session), "message": f"已写入 {len(enabled)} 个{'平差' if mode == 'adjusted' else '观测'}点的测量值", "applied_count": len(enabled), "skipped_count": len(body.mappings) - len(enabled)}

def dataset(project, include_rows):
    service=m.service_for(project)
    reader=m.MonitoringReader(service)
    catalog=reader.list_measurements()
    entry_sets = list(dict.fromkeys([*[item['key'] for item in catalog],
        *project.payload.get('axial_parameters', {}), *project.payload.get('tilt_parameters', {}),
        *(['water'] if project.payload.get('water_parameters') else []),
        *['inclinometer:' + str(p['点号']) for p in project.payload.get('measurements', {}).get('inclinometer_points', []) if p.get('点号')]]))
    return {"revision":project.revision,"entry_sets":entry_sets,"measurements":project.payload.get("measurements",{}) if include_rows else {},
        "catalog":[{**item,"dates":reader.list_dates(item["key"]),"points":reader.list_points(item["key"])} for item in catalog],
        "types":[{"key":t.key,"name":t.name,"description":t.description,"columns":m.TABLE_HEADERS[t.key]} for t in m.SINGLE_PERIOD_TYPES]}

@router.get('/catalog')
def catalog(project: CatalogProject):
    return dataset(project, False)

@router.get('/data')
def data(project: EntryProject):
    return dataset(project, True)

@router.put("/data")
def save_data(body:DataWrite,project:EntryProject,session:SessionDep):
    try:
        rows=m.normalize_rows(body.measurement_key,body.rows)
        m.validate_rows(body.measurement_key,rows)
        payload=m.save_web_rows(project,body.measurement_key,rows)
    except (ValueError,KeyError,TypeError) as exc:
        data_error(exc)
    save_project(session,project,body.revision,payload=payload)
    return public(project,session)

@router.put('/database')
def save_database(body: DatabaseWrite, project: EntryProject, session: SessionDep):
    try:
        rows = replace_date_rows(project, body.measurement_key, body.date, body.rows)
        m.validate_rows(body.measurement_key, rows)
        payload = m.save_web_rows(project, body.measurement_key, rows)
    except (ValueError, KeyError, TypeError) as exc:
        data_error(exc)
    save_project(session, project, body.revision, payload=payload)
    return public(project, session)

@router.get("/export")
def export_project(project:EntryProject):
    payload=copy.deepcopy(project.payload)
    payload["project"]={"name":project.name,"created_at":str(project.created_at),"updated_at":str(project.updated_at)}
    return Response(json.dumps(payload,ensure_ascii=False,allow_nan=False).encode("utf-8"),media_type="application/json",headers={"Content-Disposition":"attachment; filename*=UTF-8''"+quote(project.name+".proj")})

@router.post("/import")
def import_project(project:EditProject,session:SessionDep,revision:int,file:UploadFile=File(...)):
    content=bounded_upload(file)
    try:
        payload=json.loads(content.decode("utf-8-sig"),parse_constant=lambda v: (_ for _ in ()).throw(ValueError("非有限数值")))
        if not isinstance(payload,dict) or not isinstance(payload.get("measurements"),dict):
            raise ValueError("不是有效的桌面 .proj 工程文件")
        if len(payload["measurements"])>200:
            raise ValueError("单工程最多 200 个测项")
        for key,rows in payload["measurements"].items():
            if key=="inclinometer_points":
                if not isinstance(rows,list) or len(rows)>20000 or any(not isinstance(row,dict) for row in rows):
                    raise ValueError("测斜点位配置必须是表格行列表")
                continue
            rows=m.normalize_rows(key,rows)
            m.validate_rows(key,rows)
            payload["measurements"][key]=rows
        allowed={"version","measurements","axial_parameters","tilt_parameters","water_parameters","monitor_type_order"}
        clean={k:v for k,v in payload.items() if k in allowed}
        expected_types={"axial_parameters":dict,"tilt_parameters":dict,"water_parameters":list,"monitor_type_order":list}
        for name, expected in expected_types.items():
            if name in clean and not isinstance(clean[name],expected):
                raise ValueError(f"{name} 配置格式无效")
        for name in ("axial_parameters","tilt_parameters"):
            if any(not isinstance(rows,list) or any(not isinstance(row,dict) for row in rows)
                   for rows in clean.get(name,{}).values()):
                raise ValueError(f"{name} 的测项参数必须是表格行列表")
        if any(not isinstance(row,dict) for row in clean.get("water_parameters",[])):
            raise ValueError("水位参数必须是表格行列表")
        settings_in=payload.get("report_settings",{})
        if not isinstance(settings_in,dict) or not isinstance(settings_in.get("alerts",{}),dict):
            raise ValueError("报警配置格式无效")
        clean["report_settings"]={"alerts":settings_in.get("alerts",{})}
        clean["project"]={"name":project.name}
        # Imported payload must be consumable before replacing the previous data.
        candidate=project.model_copy(update={"payload":clean})
        reader=m.MonitoringReader(m.service_for(candidate))
        for item in reader.list_measurements():
            reader.list_dates(item["key"])
            reader.list_points(item["key"])
    except (ValueError,KeyError,TypeError,AttributeError,UnicodeError) as exc:
        data_error(exc)
    save_project(session,project,revision,payload=clean)
    return {**public(project,session),"message":"原始数据和计算参数已导入；桌面排版及美化数据不参与网页分析"}

@router.get("/documents")
def documents(project:AnalysisProject):
    return {"documents":[document_public(d) for d in m.knowledge_for(project).list_documents()],
            "model_configured":bool(settings.AI_BASE_URL and settings.AI_MODEL),"model":settings.AI_MODEL}

@router.post("/documents")
def add_document(project:EditProject,file:UploadFile=File(...)):
    content=bounded_upload(file)
    filename=Path((file.filename or "document").replace("\\","/")).name
    if Path(filename).suffix.lower() not in {".txt",".md",".pdf",".docx",".xlsx"}:
        raise HTTPException(422,"支持 TXT、MD、PDF、DOCX、XLSX")
    try:
        with tempfile.TemporaryDirectory(dir=settings.STORAGE_DIR) as temporary:
            path=Path(temporary)/filename
            path.write_bytes(content)
            document=m.knowledge_for(project).add_document(path)
        return document_public(document)
    except (ValueError,RuntimeError,KeyError) as exc:
        data_error(exc)

@router.delete("/documents/{document_id}")
def delete_document(document_id:str,project:EditProject):
    try:
        m.knowledge_for(project).delete_document(document_id)
    except KeyError:
        raise HTTPException(404,"资料不存在")
    return {"message":"资料已删除"}

@router.post("/analysis")
def analysis(body:Scope,project:AnalysisProject):
    try:
        return m.analysis_for(project,body)
    except (ValueError,KeyError,TypeError) as exc:
        data_error(exc)

@router.post("/ask")
def ask(body:Ask,project:AnalysisProject):
    store=m.knowledge_for(project)
    if body.mode=="documents" and not any(d["chunk_count"] for d in store.list_documents()):
        raise HTTPException(422,"请先上传并成功解析项目资料")
    client=m.model_client()
    try:
        assistant=m.AnalysisAssistant(store,m.MonitoringReader(m.service_for(project)),client)
        result=assistant.answer(body.question,body.model_dump(exclude={"question"}))
        if result.get("analysis"):
            result["analysis"]["project"]={"id":str(project.id),"name":project.name}
        return result
    except (ValueError,KeyError,TypeError) as exc:
        data_error(exc)
    except (RuntimeError,TimeoutError) as exc:
        raise HTTPException(502,str(exc)) from None

def report_for(project,scope):
    analysis=m.analysis_for(project,scope)
    return {"project_name":project.name,"scope":scope.model_dump(),"analysis":analysis,"summary_text":m.format_analysis_summary(analysis),
            "notice":"本报告采用原始监测数据与工程配置；请复核缺失数据、阈值和分析结论后签发。"}

@router.post("/reports/preview")
def preview(body:Scope,project:ReportProject):
    try:
        return report_for(project,body)
    except (ValueError,KeyError,TypeError) as exc:
        data_error(exc)

def text_cell(value):
    # Workbook contents are data, never executable Excel formulas.
    if isinstance(value,str) and value.lstrip().startswith(("=","+","-","@")):
        return "'"+value
    return value if value is not None else ""

@router.post("/reports/download")
def download(body:Scope,project:ReportProject):
    try:
        report=report_for(project,body)
    except (ValueError,KeyError,TypeError) as exc:
        data_error(exc)
    workbook=Workbook()
    sheet=workbook.active
    sheet.title="监测日报"
    sheet.append(["监测日报",text_cell(project.name)])
    sheet.append(["日期范围",body.start_date+" 至 "+body.end_date])
    sheet.append(["测项",text_cell(body.measurement_key)])
    sheet.append(["数据来源","原始监测数据"])
    sheet.append(["复核提示",report["notice"]])
    for line in report["summary_text"].splitlines():
        sheet.append([text_cell(line)])
    detail=workbook.create_sheet("分析明细")
    analysis=report["analysis"]
    detail.append(analysis["columns"])
    for row in analysis["rows"]:
        detail.append([text_cell(row.get(c)) for c in analysis["columns"]])
    warnings=workbook.create_sheet("待复核")
    warnings.append(["待复核事项"])
    for warning in analysis["warnings"]:
        warnings.append([text_cell(warning)])
    for tab in workbook:
        tab.freeze_panes="A2"
        tab.auto_filter.ref=tab.dimensions
        for cell in tab[1]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="165D8D")
        for column in tab.columns:
            tab.column_dimensions[column[0].column_letter].width=30
        for row in tab:
            for cell in row:
                cell.alignment=Alignment(vertical="top",wrap_text=True)
    output=io.BytesIO()
    workbook.save(output)
    return Response(output.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":"attachment; filename*=UTF-8''"+quote(project.name+"-监测日报.xlsx")})
