import io
import pytest
from openpyxl import Workbook
from test_projects import new_project, P


def project_url(client, user):
    return P + '/' + new_project(client, user)['id']


def save(client, user, url, key, date, rows, status=200):
    revision = client.get(url, headers=user).json()['revision']
    result = client.put(url + '/entry', headers=user, json={
        'revision': revision, 'measurement_key': key, 'date': date, 'rows': rows})
    assert result.status_code == status, result.text
    return result


def parameters(client, user, url, key, rows, status=200):
    revision = client.get(url, headers=user).json()['revision']
    result = client.put(url + '/entry/parameters', headers=user, json={
        'revision': revision, 'measurement_key': key, 'rows': rows})
    assert result.status_code == status, result.text
    return result


def entry(client, user, url, key, date):
    result = client.get(url + '/entry', headers=user, params={'measurement_key': key, 'date': date})
    assert result.status_code == 200, result.text
    return result.json()


@pytest.mark.parametrize('key, values', [('level:建筑物位移', {'测量值': 10}), ('horizontal', {'X 坐标': 10, 'Y 坐标': 20})])
def test_period_save_preserves_other_dates_and_intensified_batch(client, accounts, key, values):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    for date in ['2026-01-01', '2026-01-01（加密）', '2026-01-03']:
        save(client, user, url, key, date, [{'点号': 'A', **values}])
    assert len(client.get(url + '/data', headers=user).json()['measurements'][key]) == 3
    assert len(entry(client, user, url, key, '2026-01-03')['rows']) == 1
    save(client, user, url, key, '2026-01-03', [])
    assert len(client.get(url + '/data', headers=user).json()['measurements'][key]) == 2


def test_water_parameters_seed_readonly_columns_and_recalculate(client, accounts):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    parameters(client, user, url, 'water', [{'点位': 'W1', '管口高程（m）': 10}])
    draft = entry(client, user, url, 'water', '2026-01-01')
    assert draft['rows'][0]['管口高程（m）'] == 10
    assert {'时间', '点位', '管口高程（m）'} <= set(draft['readonly'])
    save(client, user, url, 'water', '2026-01-01', [{**draft['rows'][0], '管口高程（m）': 999, '水位读数（m）': 2}])
    assert entry(client, user, url, 'water', '2026-01-01')['rows'][0]['管口高程（m）'] == 10
    parameters(client, user, url, 'water', [{'点位': 'W1', '管口高程（m）': 12}])
    result = client.post(url + '/analysis', headers=user, json={'measurement_key': 'water', 'start_date': '2026-01-01', 'end_date': '2026-01-01'})
    assert result.status_code == 200, result.text
    assert result.json()['rows'][0]['本次水位高程(m)'] == 10
    parameters(client, user, url, 'water', [{'点位': 'W1', '管口高程（m）': 'NaN'}], 422)


def test_axial_parameters_compute_frequency_and_preview_does_not_write(client, accounts):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    key = 'axial:第一道支撑'
    parameters(client, user, url, key, [{'点号': 'Z1', '编号': '1', '传感器类型': '其他传感器',
        '标定系数 K（kN/Hz²）': .001, '初始频率 f0（Hz）': 1000, '轴力转换系数 C（无量纲）': 2}])
    draft = entry(client, user, url, key, '2026-01-01')
    assert draft['parameters'][0]['编号'] == '1'
    rows = [{'点号': 'Z1', '编号': '1', '频率': 1100, '轴力（kN）': 999}]
    revision = client.get(url, headers=user).json()['revision']
    preview = client.post(url + '/entry/preview', headers=user, json={'measurement_key': key, 'date': '2026-01-01', 'rows': rows})
    assert preview.status_code == 200, preview.text
    assert preview.json()['rows'][0]['轴力（kN）'] == 420
    assert client.get(url, headers=user).json()['revision'] == revision
    save(client, user, url, key, '2026-01-01', rows)
    assert entry(client, user, url, key, '2026-01-01')['rows'][0]['轴力（kN）'] == 420
    save(client, user, url, key, '2026-01-03', [{'点号': 'Z1', '编号': 'unconfigured', '频率': 1100}], 422)


def test_rebar_group_uses_available_frequencies_like_desktop(client, accounts):
    from app import monitoring as m
    from monitoring_daily.services.axial_force import calculate_rebar_point_force
    _, (_, user), _ = accounts
    url = project_url(client, user)
    key = 'axial:第一道支撑'
    params = [{'点号': 'Z1', '编号': str(i), '传感器类型': '钢筋计',
        '标定系数 K（kN/Hz²）': .001, '初始频率 f0（Hz）': 1000,
        '钢筋直径 d（mm）': 20, '钢筋根数 n（根）': 8, '支撑长度 L（mm）': 600, '支撑宽度 B（mm）': 800} for i in (1, 2)]
    parameters(client, user, url, key, params)
    save(client, user, url, key, '2026-01-01', [{'点号': 'Z1', '编号': str(i), '频率': 900 if i == 1 else ''} for i in (1, 2)])
    rows = entry(client, user, url, key, '2026-01-01')['rows']
    assert all(row['轴力（kN）'] == calculate_rebar_point_force({'1': 900}, params) for row in rows)


def test_tilt_links_level_points_and_updates_on_level_edit(client, accounts):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    for date, a, b in [('2026-01-01', 10, 20), ('2026-01-03', 9.99, 19.995)]:
        save(client, user, url, 'level:建筑物位移', date, [{'点号': 'A', '测量值': a}, {'点号': 'B', '测量值': b}])
    parameters(client, user, url, 'tilt:建筑物位移', [{'点号1': 'A', '点号2': 'B', '两点距离L（m）': 5}])
    draft = entry(client, user, url, 'tilt:建筑物位移', '2026-01-03')
    assert draft['rows'][0]['倾斜度（‰）'] == 1
    assert set(draft['readonly']) == set(draft['columns'])
    save(client, user, url, 'level:建筑物位移', '2026-01-03', [{'点号': 'A', '测量值': 9.98}, {'点号': 'B', '测量值': 19.995}])
    assert entry(client, user, url, 'tilt:建筑物位移', '2026-01-03')['rows'][0]['倾斜度（‰）'] == 3
    save(client, user, url, 'tilt:建筑物位移', '2026-01-03', [{'倾斜度（‰）': 999}])
    assert entry(client, user, url, 'tilt:建筑物位移', '2026-01-03')['rows'][0]['倾斜度（‰）'] == 3
    parameters(client, user, url, 'tilt:建筑物位移', [{'点号1': 'A', '点号2': 'X', '两点距离L（m）': 5}], 422)


def test_inclinometer_profile_generates_depths_and_stores_date_columns(client, accounts):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    key = 'inclinometer:CX01'
    parameters(client, user, url, key, [{'点号': 'CX01', '起始位置': '孔底', '步长': .5, '深度': '2m'}])
    draft = entry(client, user, url, key, '2026-01-01')
    assert draft['columns'] == ['深度', '测值']
    assert [r['深度'] for r in draft['rows']] == [.5, 1, 1.5, 2]
    save(client, user, url, key, '2026-01-01', [{**r, '测值': 0} for r in draft['rows']])
    save(client, user, url, key, '2026-01-03', [{**r, '测值': r['深度']} for r in draft['rows']])
    stored = client.get(url + '/data', headers=user).json()['measurements'][key]
    assert stored[-1] == {'深度': 2, '2026-01-01': 0, '2026-01-03': 2}
    parameters(client, user, url, key, [{'点号': 'CX01', '起始位置': '孔底', '步长': 1, '深度': 2}])
    adjusted = entry(client, user, url, key, '2026-01-03')['rows']
    assert adjusted == [{'深度': 1, '测值': 1}, {'深度': 2, '测值': 2}]
    save(client, user, url, key, '2026-01-03', [])
    assert '2026-01-01' in client.get(url + '/data', headers=user).json()['measurements'][key][0]
    save(client, user, url, key, '2026-01-01', [])
    assert key not in client.get(url + '/data', headers=user).json()['measurements']


def test_excel_import_is_preview_only_and_uses_desktop_sheet_mapping(client, accounts):
    _, (_, user), (_, outsider) = accounts
    url = project_url(client, user)
    book = Workbook()
    book.active.title = '单期水准数据'
    book.active.append(['序号', '时间', '点号', '测量值', '备注'])
    book.active.append([1, '2026-01-01', 'A', 10, '测试'])
    output = io.BytesIO(); book.save(output)
    params = {'measurement_key': 'level:建筑物位移', 'date': '2026-01-01'}
    files = {'file': ('entry.xlsx', output.getvalue())}
    result = client.post(url + '/entry/import', headers=user, params=params, files=files)
    assert result.status_code == 200, result.text
    assert result.json()['rows'][0]['测量值'] == 10
    assert client.get(url, headers=user).json()['revision'] == 1
    assert client.get(url + '/data', headers=user).json()['measurements'] == {}
    assert client.post(url + '/entry/import', headers=outsider, params=params, files=files).status_code == 404


def test_entry_permissions_invalid_dates_and_stale_saves(client, accounts):
    admin, (owner, user), _ = accounts
    url = project_url(client, user)
    save(client, user, url, 'level:沉降', 'bad', [{'点号': 'A', '测量值': 10}], 422)
    save(client, user, url, 'level:沉降', '2026-01-01', [{'点号': 'A', '测量值': 10}])
    assert client.put(url + '/entry', headers=user, json={'revision': 1, 'measurement_key': 'level:沉降', 'date': '2026-01-01', 'rows': []}).status_code == 409
    assert client.patch('/api/v1/users/' + owner['id'], headers=admin, json={'permissions': ['database_data']}).status_code == 200
    assert client.get(url + '/entry', headers=user, params={'measurement_key': 'water', 'date': '2026-01-01'}).status_code == 403
    assert client.put(url + '/entry/parameters', headers=user, json={'revision': 2, 'measurement_key': 'water', 'rows': []}).status_code == 403


def test_database_edit_recalculates_linked_tilt(client, accounts):
    _, (_, user), _ = accounts
    url = project_url(client, user)
    for date in ['2026-01-01', '2026-01-03']:
        save(client, user, url, 'level:沉降', date, [{'点号': 'A', '测量值': 10}, {'点号': 'B', '测量值': 20}])
    parameters(client, user, url, 'tilt:沉降', [{'点号1': 'A', '点号2': 'B', '两点距离L（m）': 5}])
    revision = client.get(url, headers=user).json()['revision']
    result = client.put(url + '/database', headers=user, json={'revision': revision, 'measurement_key': 'level:沉降', 'date': '2026-01-03', 'rows': [{'时间': '2026-01-03', '点号': 'A', '测量值': 9.99}, {'时间': '2026-01-03', '点号': 'B', '测量值': 20}]})
    assert result.status_code == 200, result.text
    stored = client.get(url + '/data', headers=user).json()['measurements']['tilt:沉降']
    assert stored[-1]['倾斜度（‰）'] == 2


def test_excel_template_and_invalid_workbook(client, accounts):
    from openpyxl import load_workbook
    _, (_, user), _ = accounts
    url = project_url(client, user)
    result = client.get(url + '/entry/template', headers=user, params={'measurement_key': 'horizontal'})
    assert result.status_code == 200, result.text
    book = load_workbook(io.BytesIO(result.content))
    assert book.sheetnames == ['单期水平位移数据']
    assert list(book.active.values)[0] == ('序号', '时间', '点号', 'X 坐标', 'Y 坐标', '备注')
    result = client.post(url + '/entry/import', headers=user, params={'measurement_key': 'horizontal', 'date': '2026-01-01'}, files={'file': ('broken.xlsx', b'broken')})
    assert result.status_code == 422
