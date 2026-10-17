import io
import json
from openpyxl import load_workbook
from test_projects import new_project, P

ROWS = [{"点号":"JGC-01","时间":"2026-01-01","测量值":10},{"点号":"JGC-01","时间":"2026-01-03","测量值":9.99}]
SCOPE = {"measurement_key":"level:建筑物位移","start_date":"2026-01-01","end_date":"2026-01-03","point":""}

def test_data_analysis_report_and_revision_conflict(client, accounts):
    _, (_, a), _ = accounts
    p = new_project(client,a)
    url=P+"/"+p["id"]
    body={"revision":1,"measurement_key":SCOPE["measurement_key"],"rows":ROWS}
    result=client.put(url+"/data",headers=a,json=body)
    assert result.status_code==200, result.text
    assert result.json()["revision"]==2
    assert client.put(url+"/data",headers=a,json=body).status_code==409
    result=client.get(url+"/data",headers=a).json()
    assert result["measurements"][SCOPE["measurement_key"]]==ROWS
    analysis=client.post(url+"/analysis",headers=a,json=SCOPE)
    assert analysis.status_code==200, analysis.text
    last=analysis.json()["rows"][-1]
    assert last["本次累计变化量（mm）"]==-10
    assert last["变化速率（mm/d）"]==-5
    assert "未知" in last["监测结论"]
    assert "path" not in analysis.json()["project"]
    export=client.post(url+"/reports/download",headers=a,json=SCOPE)
    assert export.status_code==200, export.text
    workbook=load_workbook(io.BytesIO(export.content))
    assert "分析明细" in workbook.sheetnames
    assert -10 in [v for row in workbook["分析明细"].values for v in row]
    preview=client.post(url+"/reports/preview",headers=a,json=SCOPE)
    assert preview.status_code==200 and preview.json()["analysis"]["rows"][-1]["本次累计变化量（mm）"]==-10

def test_import_and_export_keep_owner_and_reject_invalid_data(client, accounts):
    _, (alice,a), (bob,b)=accounts
    p=new_project(client,a); url=P+"/"+p["id"]
    payload={"project":{"name":"桌面工程","path":"C:/sensitive/path.proj","owner_id":bob["id"]},"measurements":{SCOPE["measurement_key"]:ROWS},"version":1}
    response=client.post(url+"/import?revision=1",headers=a,files={"file":("desktop.proj",json.dumps(payload).encode(),"application/json")})
    assert response.status_code==200,response.text
    project=client.get(url,headers=a).json()
    assert project["owner_id"]==alice["id"]
    exported=client.get(url+"/export",headers=a).json()
    assert exported["measurements"][SCOPE["measurement_key"]]==ROWS
    assert "path" not in exported["project"]
    invalid=client.put(url+"/data",headers=a,json={"revision":2,"measurement_key":SCOPE["measurement_key"],"rows":[{"点号":"A","时间":"bad","测量值":"cat"}]})
    assert invalid.status_code==422

def test_documents_and_model_configuration_are_project_scoped(client, accounts):
    _,(_,a),(_,b)=accounts
    p=new_project(client,a); url=P+"/"+p["id"]
    text="监测方案。建筑物沉降需要每周监测。发生异常应复核原始数据。"
    uploaded=client.post(url+"/documents",headers=a,files={"file":("plan.txt",text.encode("utf-8"),"text/plain")})
    assert uploaded.status_code==200,uploaded.text
    assert "source_path" not in uploaded.json()
    docs=client.get(url+"/documents",headers=a).json()
    assert len(docs["documents"])==1
    assert client.get(url+"/documents",headers=b).status_code==404
    result=client.post(url+"/ask",headers=a,json={"question":"分析沉降","mode":"documents"})
    assert result.status_code==503
    assert client.delete(url+"/documents/"+docs["documents"][0]["id"],headers=a).status_code==200

def test_bad_dates_and_missing_document_scope_are_rejected(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    result=client.post(url+"/ask",headers=a,json={"question":"分析","mode":"documents"})
    assert result.status_code==422
    assert client.post(url+"/analysis",headers=a,json={**SCOPE,"start_date":"2026-02-01"}).status_code==422

def test_workbook_formula_injection_is_exported_as_text(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    rows=[{**r,"点号":"=1+1"} for r in ROWS]
    assert client.put(url+"/data",headers=a,json={"revision":1,"measurement_key":SCOPE["measurement_key"],"rows":rows}).status_code==200
    r=client.post(url+"/reports/download",headers=a,json=SCOPE)
    workbook=load_workbook(io.BytesIO(r.content))
    assert all(c.data_type!="f" for sheet in workbook for row in sheet for c in row)

def test_report_preview_preserves_exact_scope(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    body={"revision":1,"measurement_key":SCOPE["measurement_key"],"rows":ROWS+[{**r,"点号":"OTHER","测量值":80} for r in ROWS]}
    assert client.put(url+"/data",headers=a,json=body).status_code==200
    selected={**SCOPE,"point":"JGC-01"}
    preview=client.post(url+"/reports/preview",headers=a,json=selected).json()
    assert preview["scope"]==selected
    assert {r["点号"] for r in preview["analysis"]["rows"]}=={"JGC-01"}

def test_ai_uses_real_project_evidence_and_returns_provider_answer(client,accounts,monkeypatch):
    import threading
    from http.server import HTTPServer,BaseHTTPRequestHandler
    from app.core.config import settings
    captured=[]
    class Provider(BaseHTTPRequestHandler):
        def do_POST(self):
            captured.append(json.loads(self.rfile.read(int(self.headers["Content-Length"]))))
            response=json.dumps({"choices":[{"message":{"content":"累计变化量为 -10 mm [M1]","reasoning_content":"接口返回的分析摘要"},"finish_reason":"stop"}]}).encode()
            self.send_response(200);self.send_header("Content-Type","application/json");self.end_headers();self.wfile.write(response)
        def log_message(self,*args):
            pass
    server=HTTPServer(("127.0.0.1",0),Provider)
    threading.Thread(target=server.serve_forever,daemon=True).start()
    monkeypatch.setattr(settings,"AI_BASE_URL","http://127.0.0.1:1/v1")
    monkeypatch.setattr(settings,"AI_MODEL","server-model")
    monkeypatch.setattr(settings,"AI_API_KEY","server-key")
    try:
        _,(_,a),_=accounts
        p=new_project(client,a);url=P+"/"+p["id"]
        settings_response=client.put(url+"/model-settings",headers=a,json={
            "revision":1,
            "base_url":f"http://127.0.0.1:{server.server_port}/v1",
            "model":"test-provider-model",
            "api_key":"test-private-key",
            "timeout":90,
        })
        assert settings_response.status_code==200,settings_response.text
        client.put(url+"/data",headers=a,json={"revision":2,"measurement_key":SCOPE["measurement_key"],"rows":ROWS})
        result=client.post(url+"/ask",headers=a,json={**SCOPE,"question":"累计变化多少","mode":"data"})
        assert result.status_code==200,result.text
        assert result.json()["answer"]=="累计变化量为 -10 mm [M1]"
        assert "-10.0" in captured[0]["messages"][-1]["content"] or "-10" in captured[0]["messages"][-1]["content"]
        assert "test-private-key" not in result.text
        assert "path" not in result.json()["analysis"]["project"]
    finally:
        server.shutdown();server.server_close()

def test_web_edits_preserve_entered_axial_water_and_independent_tilt(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    datasets={
        "axial:第一道支撑":[{"时间":"2026-01-01","点号":"Z1","编号":"1","频率":1000,"轴力（kN）":123}],
        "tilt:建筑物位移":[{"时间":"2026-01-01","点号1":"L1","点号2":"L2","两点距离L（m）":3,"倾斜度（‰）":.1}],
        "level:建筑物位移":ROWS,
        "water:地下水":[{"时间":"2026-01-01","点位":"W1","管口高程（m）":10,"水位读数（m）":2}]
    }
    revision=1
    for key,rows in datasets.items():
        result=client.put(url+"/data",headers=a,json={"revision":revision,"measurement_key":key,"rows":rows})
        assert result.status_code==200,result.text
        revision+=1
    saved=client.get(url+"/data",headers=a).json()["measurements"]
    assert saved["axial:第一道支撑"][0]["轴力（kN）"]==123
    assert saved["tilt:建筑物位移"]==datasets["tilt:建筑物位移"]
    changed=[{**datasets["water:地下水"][0],"管口高程（m）":12}]
    assert client.put(url+"/data",headers=a,json={"revision":revision,"measurement_key":"water:地下水","rows":changed}).status_code==200
    result=client.post(url+"/analysis",headers=a,json={**SCOPE,"measurement_key":"water:地下水","end_date":"2026-01-01"})
    assert result.status_code==200,result.text
    assert result.json()["rows"][0]["本次水位高程(m)"]==10

def test_import_rejects_invalid_inclinometer_metadata(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    result=client.post(url+"/import?revision=1",headers=a,files={"file":("bad.proj",b'{"measurements":{"inclinometer_points":null}}',"application/json")})
    assert result.status_code==422
    assert client.get(P,headers=a).status_code==200

def test_legacy_desktop_water_columns_import_as_editable_values(client,accounts):
    _,(_,a),_=accounts
    p=new_project(client,a);url=P+"/"+p["id"]
    legacy={"measurements":{"water":[{"点号":"DSW-01","时间":"2026-01-01","洞口高程":"6.1","水深":"1.6"}]}}
    result=client.post(url+"/import?revision=1",headers=a,files={"file":("old.proj",json.dumps(legacy).encode(),"application/json")})
    assert result.status_code==200,result.text
    row=client.get(url+"/data",headers=a).json()["measurements"]["water"][0]
    assert row["点位"]=="DSW-01" and row["管口高程（m）"]=="6.1"
    result=client.post(url+"/analysis",headers=a,json={**SCOPE,"measurement_key":"water","end_date":"2026-01-01"})
    assert result.status_code==200,result.text
    assert result.json()["rows"][0]["本次水位高程(m)"]==4.5
